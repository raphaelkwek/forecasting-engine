"""Parsing Bloomberg ``.xlsx`` workbook exports into the same shape
``bloomberg_csv`` produces, for the open-schema (freeform) pipeline.

A Bloomberg workbook export is two sheets: ``Data`` (``Date`` plus one column
per field) and ``Metadata`` (naming what was pulled, including the security).
This mirrors ``extraction.bloomberg_csv``'s shape — ``BloombergCsvExport``,
``Date`` plus ``{label}_{field}`` columns for every field the sheet has — not
the fixed 8-column contract the retired ``ingest.bloomberg`` module produced,
since ingestion stays open-schema (every field the export carries is kept,
not just a single named one).

The workbook-opening and ``Metadata``-reading logic here is ported from
``ingest.bloomberg`` (retired once this reader replaces it — see
``docs/superpowers/plans/2026-09-17-ingestion-consolidation-and-sprint2-fixes.md``),
since that part has nothing to do with the fixed contract and was already
correct: it reads the security from ``Metadata`` rather than trusting the
filename, the exact lesson a real mislabelled export taught that module.
"""

from __future__ import annotations

import zipfile
from io import BytesIO

import openpyxl
import pandas as pd

from forecasting_engine.extraction.bloomberg_csv import DATE_COLUMN, BloombergCsvExport, label

DATA_SHEET = "Data"
METADATA_SHEET = "Metadata"


class BloombergXlsxError(ValueError):
    """A file that is not shaped like a Bloomberg workbook export."""


def read_export(filename: str, data: bytes) -> BloombergCsvExport:
    """Parse one Bloomberg ``.xlsx`` workbook export.

    Raises ``BloombergXlsxError`` if malformed. Returns the same
    ``BloombergCsvExport`` shape ``bloomberg_csv.read_export`` does, so both
    readers feed the same ``merge()``.
    """
    try:
        book = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, OSError) as exc:
        raise BloombergXlsxError(f"{filename}: not a readable .xlsx file ({exc})") from exc

    try:
        if DATA_SHEET not in book.sheetnames:
            raise BloombergXlsxError(
                f"{filename}: no {DATA_SHEET!r} sheet found (found {book.sheetnames})"
            )
        rows = list(book[DATA_SHEET].iter_rows(values_only=True))
        security = _security(book)
    finally:
        book.close()

    if not rows:
        raise BloombergXlsxError(f"{filename}: the {DATA_SHEET!r} sheet is empty")

    header = [str(cell) for cell in rows[0]]
    frame = pd.DataFrame(rows[1:], columns=header)
    frame = frame.rename(columns={frame.columns[0]: DATE_COLUMN})
    frame[DATE_COLUMN] = pd.to_datetime(frame[DATE_COLUMN], errors="coerce")

    # Read straight from cell values rather than through pandas' CSV parser,
    # so Bloomberg's "#N/A N/A" placeholder (and anything else non-numeric)
    # needs coercing to NaN explicitly here — pd.read_csv does this for the
    # CSV reader automatically via its default NA-string list, this doesn't.
    for col in frame.columns:
        if col != DATE_COLUMN:
            frame[col] = pd.to_numeric(frame[col], errors="coerce")

    frame, notes = _dedupe_dates(frame, filename)

    lbl = label(security, filename)
    fields = [c for c in frame.columns if c != DATE_COLUMN]
    frame = frame.rename(columns={field: f"{lbl}_{field}" for field in fields})

    return BloombergCsvExport(filename=filename, security=security, frame=frame, notes=notes)


def _security(book: openpyxl.Workbook) -> str:
    """The ``Security`` value from the ``Metadata`` sheet, or "" if absent.

    Read from metadata rather than trusted from the filename — the real
    export that motivated this (``ingest.bloomberg``'s docstring) was named
    for one security and contained another.
    """
    if METADATA_SHEET not in book.sheetnames:
        return ""
    for row in book[METADATA_SHEET].iter_rows(values_only=True):
        if row and str(row[0]).strip() == "Security":
            return str(row[1]).strip()
    return ""


def _dedupe_dates(frame: pd.DataFrame, name: str) -> tuple[pd.DataFrame, tuple[str, ...]]:
    """Keep the last row for any date the sheet repeats, and say so.

    A workbook's ``Data`` sheet has every field for a date on one row, so a
    repeated date is resolved once for the whole row, not per field — unlike
    ``ingest.bloomberg.dedupe_dates``, which dedupes one field's series at a
    time. The message shape mirrors that function's.
    """
    repeated = frame[DATE_COLUMN].duplicated(keep="last")
    if not repeated.any():
        return frame, ()
    dates = sorted({d.date().isoformat() for d in frame.loc[repeated, DATE_COLUMN].dropna()})
    shown = ", ".join(dates[:5]) + (f" and {len(dates) - 5} more" if len(dates) > 5 else "")
    plural = "s" if len(dates) != 1 else ""
    note = f"{name}: {len(dates)} repeated date{plural} ({shown}); the last value on each was kept"
    return frame[~repeated].reset_index(drop=True), (note,)
