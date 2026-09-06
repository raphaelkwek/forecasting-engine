"""Writing the two-tab Excel deliverable."""

import io

import openpyxl
import pandas as pd

from forecasting_engine.extraction.workbook import build


def test_the_workbook_has_both_sheets_in_order():
    bloomberg = pd.DataFrame({"Date": ["2024-01-01"], "A_Index_PX_LAST": [100.0]})
    ff = pd.DataFrame({"Date": ["2024-01-01"], "Mkt-RF": [0.1]})

    book = openpyxl.load_workbook(io.BytesIO(build(bloomberg, ff)))

    assert book.sheetnames == ["Bloomberg Data", "Fama-French Factors"]


def test_each_sheet_holds_its_own_data():
    bloomberg = pd.DataFrame({"Date": ["2024-01-01"], "A_Index_PX_LAST": [100.0]})
    ff = pd.DataFrame({"Date": ["2024-01-01"], "Mkt-RF": [0.1]})

    book = openpyxl.load_workbook(io.BytesIO(build(bloomberg, ff)))

    assert [c.value for c in book["Bloomberg Data"][1]] == ["Date", "A_Index_PX_LAST"]
    assert [c.value for c in book["Fama-French Factors"][1]] == ["Date", "Mkt-RF"]
