"""Writing the two-tab workbook download."""

import io

import openpyxl
import pandas as pd

from forecasting_engine.ingest.workbook import build


def test_the_workbook_has_both_sheets_in_order():
    signals = pd.DataFrame({"date": ["2024-01-01"], "spx_close": [100.0]})
    factors = pd.DataFrame({"Date": ["2024-01-01"], "Mkt-RF": [0.1]})

    book = openpyxl.load_workbook(io.BytesIO(build(signals, factors)))

    assert book.sheetnames == ["Signals", "Fama-French Factors"]


def test_each_sheet_holds_its_own_data():
    signals = pd.DataFrame({"date": ["2024-01-01"], "spx_close": [100.0]})
    factors = pd.DataFrame({"Date": ["2024-01-01"], "Mkt-RF": [0.1]})

    book = openpyxl.load_workbook(io.BytesIO(build(signals, factors)))

    assert [c.value for c in book["Signals"][1]] == ["date", "spx_close"]
    assert [c.value for c in book["Fama-French Factors"][1]] == ["Date", "Mkt-RF"]
